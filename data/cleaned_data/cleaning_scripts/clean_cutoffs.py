import re
import sys
from pathlib import Path
import openpyxl
import pandas as pd

# =========================
# Configuration
# =========================

# Modifiable column configuration: Only parse columns A-H
CUTOFF_COLS_RANGE = "A:H"

EXPECTED_COLUMNS = [
    "cutoff_id",
    "year",
    "college_abbrv",
    "dte_code",
    "branch_abbrv",
    "round",
    "category",
    "percentage"
]

# =========================
# Helpers
# =========================

def remove_sheet_if_exists(excel_path: Path, sheet: str):
    """Safely removes a sheet from an existing workbook if present."""
    if not excel_path.exists():
        return
    wb_out = openpyxl.load_workbook(excel_path)
    if sheet in wb_out.sheetnames and len(wb_out.sheetnames) > 1:
        del wb_out[sheet]
        wb_out.save(excel_path)
    wb_out.close()

# =========================
# Parse CLI Arguments
# =========================

CITY = sys.argv[1].lower() if len(sys.argv) > 1 else "nashik"

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent.parent.parent

INPUT_FILE = PROJECT_ROOT / "data" / "raw_data" / f"{CITY}_raw_data.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "data" / "cleaned_data" / f"{CITY}_cleaned_data.xlsx"

if not INPUT_FILE.exists():
    print(f"\n[ERROR] Raw data file not found: {INPUT_FILE}")
    sys.exit(1)

# Inspect all available sheets in workbook
wb = openpyxl.load_workbook(INPUT_FILE, read_only=True)
all_sheet_names = wb.sheetnames
wb.close()

raw_years_input = sys.argv[2:] if len(sys.argv) > 2 else ["all"]

parsed_year_args = []
for arg in raw_years_input:
    for part in arg.split(","):
        cleaned_part = part.strip()
        if cleaned_part:
            parsed_year_args.append(cleaned_part)

is_all = any(arg.lower() == "all" for arg in parsed_year_args)

target_sheets = []

if is_all:
    for name in all_sheet_names:
        if re.match(r"^cutoffs?_\d{4}$", name, re.IGNORECASE):
            target_sheets.append(name)
    if not target_sheets:
        print(f"[WARNING] No sheets matching 'cutoffs_<year>' found in {INPUT_FILE.name}")
else:
    for y in parsed_year_args:
        sheet_candidate = f"cutoffs_{y}" if not y.startswith("cutoffs_") else y
        if sheet_candidate in all_sheet_names:
            target_sheets.append(sheet_candidate)
        else:
            matched = False
            for s in all_sheet_names:
                if s.lower() == sheet_candidate.lower() or s.lower() == f"cutoff_{y}".lower():
                    target_sheets.append(s)
                    matched = True
                    break
            if not matched:
                print(f"[INFO] Sheet '{sheet_candidate}' not found in {INPUT_FILE.name}. Skipping.")

if not target_sheets:
    print(f"\n[INFO] No cutoff sheets to process for {CITY}.")
    sys.exit(0)

print(f"\nProcessing cutoff sheets for city '{CITY}': {target_sheets}")

# =========================
# Process Each Sheet
# =========================

for sheet_name in target_sheets:
    print(f"\n--- Cleaning sheet: {sheet_name} ---")

    try:
        df = pd.read_excel(
            INPUT_FILE,
            sheet_name=sheet_name,
            usecols=CUTOFF_COLS_RANGE
        )
    except Exception as e:
        print(f"  [ERROR] Failed to read sheet '{sheet_name}': {e}")
        continue

    # Remove extra unnamed columns
    df = df.loc[:, ~df.columns.astype(str).str.contains("^Unnamed")]

    # Normalize column names: lowercase and stripped
    df.columns = [str(c).strip().lower() for c in df.columns]

    initial_row_count = len(df)

    # 1. Clean Numeric Columns first (coercing non-numeric/text headers to NaN)
    for col in ["cutoff_id", "year", "dte_code", "round", "percentage"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    # Round percentage
    if "percentage" in df.columns:
        df["percentage"] = df["percentage"].round(2)

    # 2. Clean Text Columns and remove data-type artifacts (e.g. 'VARCHAR(NN)', 'INT')
    text_cols = ["college_abbrv", "branch_abbrv", "category"]
    for col in text_cols:
        if col in df.columns:
            df[col] = (
                df[col]
                .astype(str)
                .str.strip()
                .str.upper()
            )
            df[col] = df[col].replace(["NAN", "NONE", "", "NULL"], None)
            mask_schema = df[col].astype(str).str.contains(r"VARCHAR|INTEGER|FLOAT|CHAR|PRIMARY|FOREIGN", case=False, na=False)
            df.loc[mask_schema, col] = None

    # 3. Validation: Filter out placeholder rows & invalid rows
    valid_mask = (
        df["percentage"].notnull() &
        df["year"].notnull() &
        df["round"].notnull() &
        df["branch_abbrv"].notnull() &
        (df["college_abbrv"].notnull() | df["dte_code"].notnull())
    )

    df = df[valid_mask].copy()

    skipped_rows = initial_row_count - len(df)
    if skipped_rows > 0:
        print(f"  [INFO] Filtered out {skipped_rows} empty/placeholder/header rows.")

    if df.empty:
        print(f"  [INFO] Sheet '{sheet_name}' contains no valid cutoff data rows. Skipping save.")
        # Ensure any empty sheet leftover from an earlier run is removed
        remove_sheet_if_exists(OUTPUT_FILE, sheet_name)
        continue

    # 4. Check and Drop Duplicates
    dup_subset = [c for c in ["college_abbrv", "branch_abbrv", "year", "round", "category"] if c in df.columns]
    if dup_subset:
        duplicates = df[df.duplicated(subset=dup_subset, keep=False)]
        if not duplicates.empty:
            print(f"  [WARNING] {len(duplicates)} duplicate records found on {dup_subset}. Keeping first occurrence.")
            df = df.drop_duplicates(subset=dup_subset, keep="first")

    # 5. Save Cleaned Sheet to Cleaned Workbook
    file_exists = OUTPUT_FILE.exists()

    writer_args = {
        "engine": "openpyxl",
        "mode": "a" if file_exists else "w"
    }
    if file_exists:
        writer_args["if_sheet_exists"] = "replace"

    with pd.ExcelWriter(OUTPUT_FILE, **writer_args) as writer:
        df.to_excel(
            writer,
            sheet_name=sheet_name,
            index=False
        )

    print(f"  [SUCCESS] {len(df)} records saved to '{sheet_name}' in {OUTPUT_FILE.name}")

# =========================
# Final Clean-up: Purge any invalid/empty cutoff sheets in cleaned file
# =========================
if OUTPUT_FILE.exists():
    wb_out = openpyxl.load_workbook(OUTPUT_FILE)
    sheets_to_remove = []
    for s_name in wb_out.sheetnames:
        if re.match(r"^cutoffs?_\d{4}$", s_name, re.IGNORECASE):
            ws = wb_out[s_name]
            if ws.max_row <= 1:
                sheets_to_remove.append(s_name)
            elif ws.max_row == 2:
                # If row 2 is schema description like VARCHAR
                cell_val = str(ws.cell(row=2, column=4).value or "") + str(ws.cell(row=2, column=6).value or "")
                if "VARCHAR" in cell_val:
                    sheets_to_remove.append(s_name)

    for s_name in sheets_to_remove:
        if len(wb_out.sheetnames) > 1:
            del wb_out[s_name]
            print(f"  [CLEANUP] Removed empty sheet '{s_name}' from {OUTPUT_FILE.name}")
    wb_out.save(OUTPUT_FILE)
    wb_out.close()

print("\nCutoff cleaning process finished.")
